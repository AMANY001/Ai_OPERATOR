"""Safe, idempotent workbook writes scoped to the configured output directory."""
from __future__ import annotations
import os
import shutil
from datetime import datetime
from pathlib import Path
from typing import Any
from urllib.parse import urlsplit, urlunsplit, parse_qsl, urlencode
from openpyxl import Workbook, load_workbook
from config import ensure_directories, get_settings

DEFAULT_SHEET = "Opportunities"
DEFAULT_COLUMNS = [
    "title", "company", "category", "location", "eligibility", "eligibility_status",
    "deadline", "fit_score", "ranking_reasons", "link", "description",
    "source_mode", "verification_status", "is_sample", "contact_email", "saved_at"
]


def safe_output_path(filepath: str = "opportunities.xlsx") -> Path:
    settings = get_settings()
    ensure_directories()
    candidate = Path(filepath).expanduser()
    if not candidate.is_absolute():
        normalized = str(candidate).replace("\\", "/")
        if normalized.startswith("output/"):
            normalized = normalized[len("output/"):]
        candidate = settings.output_dir / normalized
    candidate = candidate.resolve()
    if candidate != settings.output_dir.resolve() and settings.output_dir.resolve() not in candidate.parents:
        raise ValueError("Workbook path must stay inside the configured output directory.")
    if candidate.suffix.lower() != ".xlsx":
        raise ValueError("Workbook must use the .xlsx extension.")
    candidate.parent.mkdir(parents=True, exist_ok=True)
    return candidate


def _canonical_url(value: Any) -> str:
    """Normalize URL host/tracking noise without lowercasing case-sensitive paths."""
    raw = str(value or "").strip()
    if not raw:
        return ""
    try:
        parts = urlsplit(raw)
        if not parts.scheme or not parts.netloc:
            return raw.rstrip("/")
        kept = []
        for key, val in parse_qsl(parts.query, keep_blank_values=True):
            k = key.lower()
            if k.startswith("utm_") or k in {"fbclid", "gclid", "mc_cid", "mc_eid"}:
                continue
            kept.append((key, val))
        path = parts.path.rstrip("/") or "/"
        return urlunsplit((parts.scheme.lower(), parts.netloc.lower(), path, urlencode(sorted(kept)), ""))
    except Exception:
        return raw.rstrip("/")


def read_sheet_impl(filepath: str = "opportunities.xlsx", sheet_name: str = DEFAULT_SHEET) -> dict:
    path = safe_output_path(filepath)
    if not path.exists():
        return {"path": str(path), "sheet_name": sheet_name, "rows": [], "count": 0, "exists": False}
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet_name not in wb.sheetnames:
            raise ValueError(f"Worksheet '{sheet_name}' was not found. Available sheets: {wb.sheetnames}")
        ws = wb[sheet_name]
        values = ws.iter_rows(values_only=True)
        headers = next(values, ())
        header_list = [str(v) if v is not None else "" for v in headers]
        rows = []
        for row in values:
            if not any(v is not None for v in row):
                continue
            rows.append({header_list[i]: row[i] if i < len(row) else None for i in range(len(header_list)) if header_list[i]})
        return {"path": str(path), "sheet_name": sheet_name, "rows": rows, "count": len(rows), "exists": True}
    finally:
        wb.close()


def save_to_sheet_impl(rows: list[dict], filepath: str = "opportunities.xlsx", sheet_name: str = DEFAULT_SHEET) -> dict:
    if not isinstance(rows, list) or any(not isinstance(row, dict) for row in rows):
        raise ValueError("rows must be a list of JSON objects")
    path = safe_output_path(filepath)
    if os.getenv("SIMULATE_SHEET_FAILURE", "0").strip() == "1":
        raise RuntimeError("SIMULATED_SHEET_FAILURE: controlled failure before workbook modification.")
    if path.exists():
        backup_dir = get_settings().output_dir / "backups"
        backup_dir.mkdir(parents=True, exist_ok=True)
        backup_path = backup_dir / f"{path.stem}_{datetime.now().strftime('%Y%m%d_%H%M%S_%f')}.xlsx"
        shutil.copy2(path, backup_path)

    if path.exists():
        wb = load_workbook(path)
        if sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
        else:
            ws = wb.create_sheet(sheet_name)
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = sheet_name

    try:
        # Preserve pre-existing columns and append missing expected columns.
        existing_headers = [str(c.value).strip() if c.value is not None else "" for c in ws[1]] if ws.max_row >= 1 else []
        if not any(existing_headers):
            existing_headers = []
        headers = list(existing_headers)
        for col in DEFAULT_COLUMNS:
            if col not in headers:
                headers.append(col)
        for row in rows:
            for col in row:
                if col not in headers:
                    headers.append(col)
        if headers:
            for idx, header in enumerate(headers, start=1):
                ws.cell(row=1, column=idx, value=header)

        url_index = headers.index("link") + 1 if "link" in headers else None
        existing_keys = set()
        company_index = headers.index("company") + 1 if "company" in headers else None
        title_index = headers.index("title") + 1 if "title" in headers else None
        for r in range(2, ws.max_row + 1):
            key = _canonical_url(ws.cell(r, url_index).value) if url_index else ""
            if key:
                existing_keys.add(key)
            elif company_index and title_index:
                company_value = str(ws.cell(r, company_index).value or "").strip().lower()
                title_value = str(ws.cell(r, title_index).value or "").strip().lower()
                if company_value and title_value:
                    existing_keys.add((company_value, title_value))
        saved, skipped = [], []
        now = datetime.now().astimezone().isoformat(timespec="seconds")
        for raw in rows:
            row = dict(raw)
            link_key = _canonical_url(row.get("link"))
            fallback_key = (str(row.get("company", "")).strip().lower(), str(row.get("title", "")).strip().lower())
            if (link_key and link_key in existing_keys) or (not link_key and fallback_key in existing_keys):
                skipped.append(row.get("link") or row.get("title"))
                continue
            row.setdefault("saved_at", now)
            values = []
            for header in headers:
                value = row.get(header, "")
                if isinstance(value, (dict, list)):
                    value = "; ".join(map(str, value)) if isinstance(value, list) else str(value)
                values.append(value)
            ws.append(values)
            if link_key:
                existing_keys.add(link_key)
            elif fallback_key:
                existing_keys.add(fallback_key)
            saved.append(row.get("link") or row.get("title"))
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for col_cells in ws.columns:
            max_len = min(max((len(str(cell.value or "")) for cell in col_cells), default=10), 55)
            ws.column_dimensions[col_cells[0].column_letter].width = max(12, max_len + 2)
        wb.save(path)
    finally:
        wb.close()
    checked = read_sheet_impl(filepath, sheet_name)
    return {"path": str(path), "sheet_name": sheet_name, "requested": len(rows),
            "saved": len(saved), "skipped_duplicates": len(skipped), "saved_keys": saved,
            "skipped_keys": skipped, "rows_after_write": checked["count"], "exists": checked["exists"]}
